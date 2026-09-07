"""Rebuild both CVs from the pristine originals.

The originals live in #Work/CV_backup_*/ and have exactly 30 paragraphs.
Every index below is absolute against that 30-paragraph source, so the
script MUST be run against a restored original — never against its own
output. The assert in rewrite() enforces that; without it, a second run
lands the edits on shifted paragraphs and silently duplicates entries
and drops section headers.

Usage:
    1. Copy the four files from #Work/CV_backup_*/ over #Work/
    2. python tools/rewrite_cv.py
    3. Convert the two .docx to PDF with Word (see README)
"""
import copy
import os
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Emu, Inches, Pt

WORK = r'C:\Users\jaspe\Documents' + '\\' + '#Work'
OUT_DIR = os.environ.get('CV_OUT', WORK)


def set_para(p, segments):
    """Replace a paragraph's runs with (text, bold) segments, cloning the
    first run's rPr so the eastAsia font survives for Chinese."""
    tmpl = p.runs[0]._element.get_or_add_rPr() if p.runs else None
    tmpl = copy.deepcopy(tmpl) if tmpl is not None else None
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, bold in segments:
        run = p.add_run(text)
        if tmpl is not None:
            run._element.insert(0, copy.deepcopy(tmpl))
        run.bold = bold


def clone_after(p, segments):
    el = copy.deepcopy(p._element)
    p._element.addnext(el)
    set_para(Paragraph(el, p._parent), segments)


def drop(p):
    p._element.getparent().remove(p._element)


def tighten(d):
    """Halve the 12pt header spacing; buys ~3 lines toward the one-page goal."""
    n = 0
    for p in d.paragraphs:
        sa = p.paragraph_format.space_after
        if sa is not None and sa > Emu(76200):
            p.paragraph_format.space_after = Emu(76200)
            n += 1
    return n


HEADERS = {'EDUCATION', 'SKILLS', 'PROJECTS / EXPERIENCE', 'ACHIEVEMENT',
           '教育背景', '技能', '项目 / 经历', '获奖与成就'}


def resize(d):
    """12pt body -> 11pt so everything fits one page; name and headers scale with it."""
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        size = Pt(17) if i == 0 else Pt(13) if t in HEADERS else Pt(11)
        for r in p.runs:
            r.font.size = size


def set_margins(d, inches=0.4):
    for sec in d.sections:
        sec.top_margin = Inches(inches)
        sec.bottom_margin = Inches(inches)


CONTACT = [('[+86] 15691728981 | linkedin.com/in/jasper-tan-xian | jteh37449.github.io', False)]

EN = {
    2: CONTACT,
    5: [('Activities: ', True),
        ('PPI Xi\u2019an (Indonesian Student Association) Committee 09/23~07/24; '
         'XJTU Dragon & Lion Dance Association 02/24~07/24; '
         'XJTU PSA English Association 09/24~Present', False)],
    14: [('Context: ', True),
         ('a customer-service digital employee (CS-DE) needed third-party ticket integration '
          'and lower cost per conversation. ', False),
         ('Action: ', True),
         ('added region-scoped board data fetching, ticket create/update and PDF / DOCX drafting '
          'from chat; profiled token usage and migrated the agent across LLMs with adversarial '
          'checks for falsely reported success. ', False),
         ('Result: ', True),
         ('cut model cost to ', False), ('8% of baseline (11.8x cheaper)', True),
         (', holding correctness at 26/27 against a 27/27 baseline; rejected a cheaper model '
          'that failed the write-authorisation guard.', False)],
    15: [('Kaggle x Google Competition Chess Engine Design (', False), ('12/24 ~ 02/25', True), (')', False)],
    16: [('Context: ', True),
         ('engine capped at 5 MiB RAM, one 2.20 GHz core and a 64 KB submission. ', False),
         ('Action: ', True),
         ('benchmarked open-source engines and optimised one to fit without losing playing strength. ', False),
         ('Result: ', True),
         ('18th of 1120 teams (top 2%), Kaggle Competition Silver Medal.', False)],
    17: [('MERN Stack Chat App (', False), ('05/25 ~ 06/25', True), (')', False)],
    18: [('Context: ', True),
         ('needed message delivery without polling. ', False),
         ('Action: ', True),
         ('built a full-stack app on MongoDB / Express / React / Node.js with Socket.IO '
          'bi-directional transport. ', False),
         ('Result: ', True),
         ('released open-source on GitHub.', False)],
    23: [('2025 ICPC Shaanxi Provincial Contest Silver Medal ', False), ('05/25', True)],
    26: [('Kaggle FIDE & Google Efficient Chess AI Challenge Silver Medal, 18 / 1120 teams (Top 2%) ', False),
         ('12/24 ~ 02/25', True)],
    27: [('China Computer Federation (CCF) NOI-Pre Problem Bank Tester, 100 hours volunteer service ', False),
         ('2024 ~ 2025', True)],
    28: [('Excellence Award, 4th XJTU 3-Minute Academic English Speech Competition ', False), ('11/24', True)],
    29: [('4th International Chinese Traditional Sports Championship \u2013 6th, Traditional Dragon Dance ', False),
         ('06/24', True), (', and more', False)],
}
EN_EXTRA = (27, [('Problem tester for Codeforces Rounds 1050 / 1074 / 1089 / 1119, '
                  'TLX TROC #43 and COMPFEST 18 SCPC ', False), ('2025 ~ 2026', True)])

ZH = {
    2: CONTACT,
    5: [('社团经历：', True),
        ('PPI Xi\u2019an（西安印度尼西亚学生协会）委员 09/23~07/24；'
         '西安交通大学舞龙舞狮协会成员 02/24~07/24；'
         '西安交通大学 PSA 英语协会成员 09/24~至今', False)],
    14: [('背景：', True),
         ('客户服务数字员工（CS-DE）需接入外部工单系统并降低单次对话成本。', False),
         ('行动：', True),
         ('实现按区域获取看板数据、通过对话创建与更新工单及生成 PDF / DOCX 草稿；分析 token 消耗，'
          '完成大模型迁移，并针对智能体谎报执行成功设计对抗性验证。', False),
         ('成果：', True),
         ('将模型成本降至基线的 ', False), ('8%（约 11.8 倍降本）', True),
         ('，集成测试正确率保持 26/27（基线 27/27）；并否决了成本更低但未通过写入授权校验的模型。', False)],
    16: [('背景：', True),
         ('赛题限定 5 MiB 内存、2.20 GHz 单核 CPU 与 64 KB 提交体积。', False),
         ('行动：', True),
         ('测试并优化多个开源引擎，使其在该预算内运行且保持棋力。', False),
         ('成果：', True),
         ('1120 支队伍中第 18 名（前 2%），获 Kaggle 竞赛银牌。', False)],
    18: [('背景：', True),
         ('需实现无需轮询的实时消息收发。', False),
         ('行动：', True),
         ('基于 MERN 技术栈与 Socket.IO 构建双向通信全栈应用。', False),
         ('成果：', True),
         ('已在 GitHub 开源。', False)],
    23: [('2025 年 ICPC 陕西省赛银牌 ', False), ('05/25', True)],
    26: [('Kaggle FIDE & Google 高效国际象棋 AI 挑战赛银牌，1120 支队伍中第 18 名（前 2%）', False),
         ('12/24 ~ 02/25', True)],
    27: [('中国计算机学会（CCF）NOI-Pre 题库测试志愿者，累计志愿服务 100 小时 ', False), ('2024 ~ 2025', True)],
    28: [('第四届西安交通大学三分钟学术英语演讲比赛优秀奖 ', False), ('11/24', True)],
    29: [('第四届中华传统体育国际锦标赛传统舞龙第 6 名 ', False), ('06/24', True), ('，及其他荣誉', False)],
}
ZH_EXTRA = (27, [('Codeforces Round 1050 / 1074 / 1089 / 1119、TLX TROC #43 及 '
                  'COMPFEST 18 SCPC 题目测试员 ', False), ('2025 ~ 2026', True)])

DROP = [6, 7]   # the two society bullets merged into paragraph 5


def rewrite(path, mapping, extra, out):
    d = Document(path)
    assert len(d.paragraphs) == 30, (
        f'{path}: expected the 30-paragraph original, got {len(d.paragraphs)}. '
        'Restore from #Work/CV_backup_* first — the indices here are absolute.')
    paras = list(d.paragraphs)              # snapshot keeps indices stable
    for idx, segs in mapping.items():
        set_para(paras[idx], segs)
    clone_after(paras[extra[0]], extra[1])
    for i in DROP:
        drop(paras[i])
    tighten(d)
    resize(d)
    set_margins(d)
    d.save(out)
    print('wrote', out, '-', len(Document(out).paragraphs), 'paragraphs')


if __name__ == '__main__':
    rewrite(os.path.join(WORK, 'Jasper Tan_EN.docx'), EN, EN_EXTRA,
            os.path.join(OUT_DIR, 'Jasper Tan_EN.docx'))
    rewrite(os.path.join(WORK, 'Jasper Tan_中文简历.docx'), ZH, ZH_EXTRA,
            os.path.join(OUT_DIR, 'Jasper Tan_中文简历.docx'))
